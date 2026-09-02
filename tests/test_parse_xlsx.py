"""Разбор xlsx: вкладки, скрытые листы, заголовок, номера строк."""

import io

import pytest
from openpyxl import Workbook

from anki_deck_gen.errors import TableUnreadable
from anki_deck_gen.tables.parse import parse_xlsx


def _workbook(
    sheets: dict[str, list[list[object]]], *, hidden: frozenset[str] = frozenset()
) -> bytes:
    workbook = Workbook()
    first = True
    for name, rows in sheets.items():
        if first:
            worksheet = workbook.active
            assert worksheet is not None
            worksheet.title = name
            first = False
        else:
            worksheet = workbook.create_sheet(name)
        for row in rows:
            worksheet.append(row)
        if name in hidden:
            worksheet.sheet_state = "hidden"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_each_visible_sheet_becomes_a_sheet_with_spreadsheet_row_numbers() -> None:
    data = _workbook(
        {
            "Greetings": [
                ["Q", "A", "Tags"],
                ["Hi", "Привет", "a, b"],
                [None, None, None],
                ["Bye", "Пока", ""],
            ],
            "Body": [["Вопрос", "Ответ"], ["hand", "рука"]],
        }
    )
    table = parse_xlsx(data, title="Book")
    assert table.title == "Book"
    assert [s.name for s in table.sheets] == ["Greetings", "Body"]
    greetings = table.sheets[0]
    assert [r.number for r in greetings.rows] == [2, 4], "the blank row is skipped, numbering kept"
    assert greetings.rows[0].tags == ("a", "b")
    assert table.sheets[1].columns == frozenset({"Q", "A"})
    assert table.multi_sheet


def test_hidden_sheets_are_skipped() -> None:
    data = _workbook(
        {"Main": [["Q", "A"], ["x", "y"]], "Scratch": [["junk"], ["more"]]},
        hidden=frozenset({"Scratch"}),
    )
    table = parse_xlsx(data)
    assert [s.name for s in table.sheets] == ["Main"]
    assert not table.multi_sheet


def test_a_header_only_sheet_is_kept_with_zero_rows() -> None:
    data = _workbook({"Main": [["Q", "A"], ["x", "y"]], "Empty": [["Q", "A"]]})
    table = parse_xlsx(data)
    assert [len(s.rows) for s in table.sheets] == [1, 0]


def test_a_sheet_without_q_and_a_names_the_sheet() -> None:
    data = _workbook(
        {"Main": [["Q", "A"], ["x", "y"]], "Oops": [["Word", "Translation"], ["a", "b"]]}
    )
    with pytest.raises(TableUnreadable) as info:
        parse_xlsx(data)
    assert "Лист «Oops»" in info.value.detail
    assert "Word | Translation" in info.value.detail


def test_deck_column_and_custom_columns_are_carried() -> None:
    data = _workbook({"S": [["Q", "A", "Колода", "Tips"], ["a", "b", "Часть 1", "hint"]]})
    row = parse_xlsx(data).rows[0]
    assert row.deck == "Часть 1"
    assert row.extra == {"Tips": "hint"}


def test_numbers_lose_the_trailing_zero() -> None:
    data = _workbook({"S": [["Q", "A"], [5.0, 2.5]]})
    row = parse_xlsx(data).rows[0]
    assert (row.question, row.answer) == ("5", "2.5")


def test_garbage_bytes_are_unreadable() -> None:
    with pytest.raises(TableUnreadable):
        parse_xlsx(b"not a workbook")


def test_a_workbook_with_no_data_sheets_is_unreadable() -> None:
    # Обе вкладки видимы, но первая строка пуста — для разбора это «вкладка без данных».
    data = _workbook({"One": [[None, None]], "Two": [[None]]})
    with pytest.raises(TableUnreadable):
        parse_xlsx(data)
